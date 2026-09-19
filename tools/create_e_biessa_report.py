from __future__ import annotations

import math
import os
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Jasser Noomani\Desktop\Stage_2026")
OUT = ROOT / "deliverables"
ASSETS = ROOT / "report_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

DOCX_PATH = OUT / "Rapport_de_stage_E-biessa.docx"
LOGO = ROOT / "marketplace-frontend" / "public" / "e-biessa-logo.png"

GREEN = "073D2C"
GREEN_2 = "0B5A40"
MINT = "DDEFE7"
PALE = "F2F7F4"
ORANGE = "FFE5DC"
ACCENT = "FF5A36"
INK = "17211D"
GRAY = "66736D"
LIGHT_GRAY = "E5ECE8"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\aptos.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibri.ttf"),
    ]
    if bold:
        candidates = [
            Path(r"C:\Windows\Fonts\aptos-bold.ttf"),
            Path(r"C:\Windows\Fonts\arialbd.ttf"),
            Path(r"C:\Windows\Fonts\calibrib.ttf"),
        ] + candidates
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered(draw, box, title, subtitle="", fill="#ffffff", outline="#_unused"):
    x1, y1, x2, y2 = box
    outline_color = "#BFD8CC" if outline == "#_unused" else outline
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline_color, width=3)
    title_font = font(28, True)
    sub_font = font(19)
    title_lines = wrap(draw, title, title_font, x2 - x1 - 44)
    sub_lines = wrap(draw, subtitle, sub_font, x2 - x1 - 44) if subtitle else []
    line_h = 34
    sub_h = 25
    total = len(title_lines) * line_h + (12 + len(sub_lines) * sub_h if sub_lines else 0)
    y = y1 + (y2 - y1 - total) / 2
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=title_font, fill="#073D2C")
        y += line_h
    if sub_lines:
        y += 12
        for line in sub_lines:
            bbox = draw.textbbox((0, 0), line, font=sub_font)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=sub_font, fill="#51635B")
            y += sub_h


def arrow(draw, start, end, color="#FF5A36", width=7):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 19
    wing = 0.55
    p1 = (end[0] - length * math.cos(angle - wing), end[1] - length * math.sin(angle - wing))
    p2 = (end[0] - length * math.cos(angle + wing), end[1] - length * math.sin(angle + wing))
    draw.polygon([end, p1, p2], fill=color)


def base_canvas(w=1600, h=900, title=""):
    img = Image.new("RGB", (w, h), "#F7FAF8")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((20, 20, w - 20, h - 20), radius=34, fill="#F7FAF8", outline="#D7E4DE", width=3)
    if title:
        d.text((70, 52), title, font=font(38, True), fill="#17211D")
        d.line((70, 112, 240, 112), fill="#FF5A36", width=8)
    return img, d


def create_diagrams():
    # Architecture globale
    img, d = base_canvas(title="Architecture globale de la solution")
    draw_centered(d, (80, 190, 390, 390), "Application Angular", "Interface responsive\nAcheteur - Vendeur - Admin", fill="#FFFFFF")
    draw_centered(d, (540, 165, 1000, 415), "API Spring Boot", "Sécurité JWT/OAuth2 - règles métier\nREST - WebSocket - synchronisation", fill="#DDEFE7", outline="#7BB49A")
    draw_centered(d, (1160, 90, 1500, 270), "MySQL", "Source transactionnelle", fill="#FFFFFF")
    draw_centered(d, (1160, 330, 1500, 510), "Stockage images", "Local ou Cloudinary", fill="#FFFFFF")
    draw_centered(d, (510, 610, 1010, 805), "Service IA FastAPI", "Compréhension de requête - embeddings\nrecherche sémantique - orchestration", fill="#FFE5DC", outline="#FF9A82")
    draw_centered(d, (80, 640, 390, 805), "Qdrant", "Index vectoriel persistant", fill="#FFFFFF")
    draw_centered(d, (1160, 610, 1500, 805), "Ollama + Qwen 2.5", "Modèle adapté par QLoRA", fill="#FFFFFF")
    arrow(d, (390, 290), (540, 290))
    arrow(d, (1000, 235), (1160, 180))
    arrow(d, (1000, 340), (1160, 415))
    arrow(d, (770, 415), (770, 610))
    arrow(d, (510, 710), (390, 720))
    arrow(d, (1010, 710), (1160, 710))
    img.save(ASSETS / "architecture.png", quality=95)

    # Modèle de données
    img, d = base_canvas(h=1020, title="Vue simplifiée du modèle de données")
    boxes = {
        "Utilisateur": (90, 170, 430, 340, "rôle, profil, statut\npréférences véhicules"),
        "Annonce": (630, 170, 970, 340, "vendeur, catégorie, prix\nétat, livraison, garantie"),
        "Véhicule": (1160, 170, 1500, 340, "marque, modèle, année\nmoteur, carburant"),
        "Image": (1160, 470, 1500, 620, "ordre, image principale"),
        "Conversation": (90, 470, 430, 620, "acheteur, vendeur, annonce"),
        "Message": (90, 750, 430, 900, "texte, pièce jointe, date"),
        "Avis vendeur": (630, 750, 970, 900, "note 1 à 5, commentaire"),
        "Signal activité": (1160, 750, 1500, 900, "vue, recherche, contact"),
        "Signalement": (630, 470, 970, 620, "motif, statut, résolution"),
    }
    fills = ["#FFFFFF", "#DDEFE7", "#FFFFFF", "#FFFFFF", "#FFE5DC", "#FFFFFF", "#FFFFFF", "#FFFFFF", "#FFF3EE"]
    for ((name, (x1, y1, x2, y2, sub)), fillc) in zip(boxes.items(), fills):
        draw_centered(d, (x1, y1, x2, y2), name, sub, fill=fillc)
    links = [
        ((430, 250), (630, 250)), ((970, 250), (1160, 250)), ((970, 300), (1280, 470)),
        ((260, 340), (260, 470)), ((260, 620), (260, 750)), ((430, 550), (630, 550)),
        ((430, 300), (700, 750)), ((430, 290), (1200, 750)), ((800, 340), (800, 470)),
    ]
    for s, e in links:
        arrow(d, s, e, color="#7A9488", width=5)
    img.save(ASSETS / "data_model.png", quality=95)

    # Recherche IA
    img, d = base_canvas(title="Chaîne de recherche intelligente et réponse ancrée")
    steps = [
        (60, 260, 300, 500, "1  Requête libre", "Français, anglais ou arabe tunisien"),
        (365, 260, 605, 500, "2  Extraction", "Qwen produit des contraintes structurées"),
        (670, 260, 910, 500, "3  Recherche", "Embeddings FastEmbed et candidats Qdrant"),
        (975, 260, 1215, 500, "4  Validation", "Compatibilité déterministe et seuil"),
        (1280, 260, 1520, 500, "5  Réponse", "Annonces actives et explications vérifiables"),
    ]
    for i, (x1, y1, x2, y2, t, s) in enumerate(steps):
        draw_centered(d, (x1, y1, x2, y2), t, s, fill="#DDEFE7" if i in (1, 3) else "#FFFFFF")
        if i < len(steps) - 1:
            arrow(d, (x2, 380), (steps[i + 1][0], 380))
    d.rounded_rectangle((220, 620, 1380, 760), radius=25, fill="#073D2C")
    msg = "Principe de sûreté : le modèle propose, le catalogue et les règles métier décident."
    bbox = d.textbbox((0, 0), msg, font=font(26, True))
    d.text(((1600 - (bbox[2] - bbox[0])) / 2, 672), msg, font=font(26, True), fill="#FFFFFF")
    img.save(ASSETS / "ai_pipeline.png", quality=95)

    # Messagerie temps réel
    img, d = base_canvas(title="Séquence simplifiée de la messagerie instantanée")
    actors = [(160, "Acheteur"), (560, "Angular"), (980, "Spring Boot"), (1380, "Vendeur")]
    for x, label in actors:
        d.rounded_rectangle((x - 120, 145, x + 120, 225), radius=18, fill="#FFFFFF", outline="#AFC9BD", width=3)
        bbox = d.textbbox((0, 0), label, font=font(24, True))
        d.text((x - (bbox[2] - bbox[0]) / 2, 168), label, font=font(24, True), fill="#073D2C")
        d.line((x, 225, x, 820), fill="#C7D7D0", width=3)
    events = [
        (285, 160, 560, "Écrit un message ou joint une image"),
        (385, 560, 980, "POST authentifié et fichier validé"),
        (500, 980, 1380, "Diffusion WebSocket avec ticket court"),
        (610, 1380, 560, "Mise à jour immédiate de la conversation"),
        (720, 560, 160, "Accusé d'affichage et défilement au dernier message"),
    ]
    for y, x1, x2, label in events:
        arrow(d, (x1, y), (x2, y), color="#FF5A36" if x2 > x1 else "#0B5A40", width=6)
        lines = wrap(d, label, font(19, True), abs(x2 - x1) - 40)
        d.text((min(x1, x2) + 20, y - 46), lines[0], font=font(19, True), fill="#17211D")
    img.save(ASSETS / "messaging_sequence.png", quality=95)

    # Recommandation
    img, d = base_canvas(h=760, title="Recommandation personnalisée explicable")
    blocks = [
        (65, 205, 360, 500, "Préférences", "Véhicules déclarés\noption de personnalisation"),
        (470, 205, 765, 500, "Signaux privés", "Vues, favoris, catégories\ncontacts récents"),
        (875, 205, 1170, 500, "Classement", "Compatibilité stricte\npondération et diversité"),
        (1280, 205, 1535, 500, "Suggestions", "Raisons lisibles\nannonces actives"),
    ]
    for i, (x1, y1, x2, y2, t, s) in enumerate(blocks):
        draw_centered(d, (x1, y1, x2, y2), t, s, fill="#FFFFFF" if i != 2 else "#DDEFE7")
        if i < len(blocks) - 1:
            arrow(d, (x2, 350), (blocks[i + 1][0], 350))
    d.text((70, 625), "Exclusions : annonces vendues, vendeurs suspendus, propres annonces et éléments déjà sauvegardés.", font=font(22, True), fill="#073D2C")
    img.save(ASSETS / "recommendations.png", quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in edges:
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key, value in edges[edge].items():
                element.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=10.5, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def field(run, instruction):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=GRAY)
    run2 = paragraph.add_run()
    set_run_font(run2, size=8.5, color=GRAY)
    field(run2, "PAGE")


def setup_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, color in [("Heading 1", 20, GREEN), ("Heading 2", 14, GREEN_2), ("Heading 3", 11.5, INK)]:
        st = doc.styles[style_name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 4)
        st.paragraph_format.space_after = Pt(6)

    if "Caption Ebiessa" not in [s.name for s in doc.styles]:
        st = doc.styles.add_style("Caption Ebiessa", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Aptos"
        st.font.size = Pt(8.5)
        st.font.italic = True
        st.font.color.rgb = RGBColor.from_string(GRAY)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.space_after = Pt(8)

    if "Lead" not in [s.name for s in doc.styles]:
        st = doc.styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Aptos"
        st.font.size = Pt(12.5)
        st.font.color.rgb = RGBColor.from_string(GREEN)
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.space_after = Pt(10)


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)


def configure_headers(section, show=True):
    header = section.header
    header.is_linked_to_previous = False
    footer = section.footer
    footer.is_linked_to_previous = False
    if show:
        p = header.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("E-biessa   |   Rapport de stage")
        set_run_font(r, size=8.5, bold=True, color=GREEN)
        p.paragraph_format.space_after = Pt(2)
        p2 = footer.paragraphs[0]
        p2.clear()
        add_page_number(p2)
    else:
        header.paragraphs[0].clear()
        footer.paragraphs[0].clear()


def add_text(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r, size=12.5 if style == "Lead" else 10.5, color=GREEN if style == "Lead" else INK)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.6 + level * 0.45)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{index}. {item}")
        set_run_font(r)


def add_callout(doc, title, body, fill=MINT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": "8", "color": GREEN_2}, bottom={"val": "single", "sz": "8", "color": GREEN_2}, left={"val": "single", "sz": "14", "color": ACCENT}, right={"val": "single", "sz": "8", "color": GREEN_2})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title + "\n")
    set_run_font(r, bold=True, color=GREEN)
    r2 = p.add_run(body)
    set_run_font(r2, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False if widths else True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(h))
        set_run_font(r, size=9, bold=True, color=WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        if ri % 2 == 1:
            for c in cells:
                set_cell_shading(c, PALE)
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(value))
            set_run_font(r, size=8.7)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, path, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph(caption, style="Caption Ebiessa")
    return cap


def add_chapter(doc, number, title, intro):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"CHAPITRE {number}")
    set_run_font(r, size=11, bold=True, color=ACCENT)
    h = doc.add_paragraph(f"{number}. {title}", style="Heading 1")
    h.paragraph_format.space_after = Pt(6)
    add_text(doc, intro, style="Lead")


def add_major_title(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        set_run_font(run, name="Aptos Display", size=25, bold=True, color=INK)
    return p


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "La table des matières sera mise à jour automatiquement à l'ouverture du document."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_char2, placeholder, fld_char3])


def add_cover(doc):
    section = doc.sections[0]
    configure_section(section)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    configure_headers(section, show=False)

    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = banner.cell(0, 0)
    set_cell_shading(cell, GREEN)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("RAPPORT DE STAGE ACADÉMIQUE")
    set_run_font(r, size=13, bold=True, color=WHITE)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(11.2))
        p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Conception et réalisation de E-biessa")
    set_run_font(r, name="Aptos Display", size=25, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Place de marché intelligente de pièces automobiles d'occasion")
    set_run_font(r, size=14, bold=True, color=GREEN_2)

    accent = doc.add_table(rows=1, cols=1)
    accent.alignment = WD_TABLE_ALIGNMENT.CENTER
    accent.cell(0, 0).width = Cm(4)
    set_cell_shading(accent.cell(0, 0), ACCENT)
    accent.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    details = [
        ("Réalisé par", "Jasser Noomani"),
        ("Domaine", "Génie logiciel et intelligence artificielle"),
        ("Nature", "Projet académique de fin de stage"),
        ("Année universitaire", "2025-2026"),
    ]
    for row, (k, v) in zip(meta.rows, details):
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(9.0)
        set_cell_shading(row.cells[0], PALE)
        set_cell_shading(row.cells[1], WHITE)
        set_cell_border(row.cells[0], bottom={"val": "single", "sz": "4", "color": LIGHT_GRAY})
        set_cell_border(row.cells[1], bottom={"val": "single", "sz": "4", "color": LIGHT_GRAY})
        r1 = row.cells[0].paragraphs[0].add_run(k)
        set_run_font(r1, size=10, bold=True, color=GREEN)
        r2 = row.cells[1].paragraphs[0].add_run(v)
        set_run_font(r2, size=10)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("Une plateforme sécurisée, explicable et centrée sur la compatibilité véhicule")
    set_run_font(r, size=10.5, italic=True, color=GRAY)


def build_document():
    create_diagrams()
    doc = Document()
    setup_styles(doc)
    add_cover(doc)

    # New section after cover, with header/footer.
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(sec)
    configure_headers(sec, show=True)

    doc.add_paragraph("Remerciements", style="Title")
    add_text(doc, "Je tiens à remercier toutes les personnes qui ont contribué, directement ou indirectement, à la réalisation de ce projet. Leurs échanges, leurs retours et leurs conseils ont permis de transformer un cahier des charges ambitieux en une solution cohérente, testable et documentée.")
    add_text(doc, "Ce travail a également été l'occasion d'approfondir des domaines complémentaires : conception d'API sécurisées, développement d'interfaces Angular, modélisation de données, recherche sémantique, adaptation d'un modèle de langage et validation de règles de compatibilité automobile.")

    doc.add_page_break()
    doc.add_paragraph("Résumé", style="Title")
    add_text(doc, "E-biessa est une place de marché intelligente dédiée aux pièces automobiles d'occasion en Tunisie. La plateforme réunit trois profils - acheteur, vendeur et administrateur - autour d'un parcours complet : création de compte, gestion du profil, publication d'annonces illustrées, recherche multicritère, contact, messagerie instantanée, modération et suivi de la réputation.")
    add_text(doc, "L'originalité de la solution réside dans une architecture hybride. Spring Boot garantit les transactions, les droits d'accès et la cohérence du catalogue. Angular fournit une interface responsive en modes clair et sombre. Un service FastAPI réalise l'indexation vectorielle et la recherche sémantique. Un modèle Qwen 2.5 7B, adapté localement par QLoRA, extrait les contraintes véhicule, carburant, position et type de pièce à partir de requêtes en français, en anglais ou en arabe tunisien. Les résultats sont ensuite contrôlés par des règles déterministes afin d'éviter toute compatibilité inventée.")
    add_text(doc, "La version finale intègre aussi des recommandations personnalisées explicables, une orchestration agentique en lecture seule et une stratégie de tests couvrant le backend, l'interface, le service d'intelligence artificielle et les parcours navigateur. Le projet démontre qu'une IA générative peut améliorer la découverte d'annonces sans devenir la source de vérité du domaine.")
    p = doc.add_paragraph()
    r = p.add_run("Mots-clés : ")
    set_run_font(r, bold=True, color=GREEN)
    r2 = p.add_run("marketplace, pièces automobiles, Spring Boot, Angular, FastAPI, Qdrant, Qwen, QLoRA, RAG, recommandation, sécurité.")
    set_run_font(r2)

    doc.add_paragraph("Abstract", style="Title")
    add_text(doc, "E-biessa is an intelligent marketplace for used automotive parts in Tunisia. It supports buyers, sellers and administrators through secure account management, rich listings, multi-vehicle fitment, image handling, moderation, seller reputation, direct contact and real-time messaging.")
    add_text(doc, "The final solution uses a hybrid architecture: Spring Boot remains the transactional and security boundary, Angular delivers the user experience, and FastAPI hosts semantic retrieval and model-backed query understanding. A locally fine-tuned Qwen 2.5 7B model extracts structured automotive constraints, while deterministic catalogue rules remain authoritative. Explainable recommendations and a read-only agent orchestration layer complete the intelligent search journey.")
    p = doc.add_paragraph()
    r = p.add_run("Keywords: ")
    set_run_font(r, bold=True, color=GREEN)
    r2 = p.add_run("automotive marketplace, semantic search, fitment, Spring Boot, Angular, FastAPI, Qdrant, QLoRA, recommendations.")
    set_run_font(r2)

    doc.add_page_break()
    doc.add_paragraph("Sommaire", style="Title")
    add_toc(doc)

    doc.add_page_break()
    doc.add_paragraph("Liste des abréviations", style="Title")
    add_table(doc, ["Abréviation", "Signification"], [
        ("API", "Application Programming Interface"),
        ("DTO", "Data Transfer Object"),
        ("E2E", "End-to-End"),
        ("JWT", "JSON Web Token"),
        ("LLM", "Large Language Model"),
        ("OEM", "Original Equipment Manufacturer"),
        ("QLoRA", "Quantized Low-Rank Adaptation"),
        ("RAG", "Retrieval-Augmented Generation"),
        ("RBAC", "Role-Based Access Control"),
        ("SPA", "Single-Page Application"),
        ("UML", "Unified Modeling Language"),
        ("WebSocket", "Canal bidirectionnel persistant entre client et serveur"),
    ], widths=[3.2, 12.5])

    doc.add_paragraph("Liste des figures", style="Heading 2")
    add_bullets(doc, [
        "Figure 1 - Architecture globale de la solution",
        "Figure 2 - Vue simplifiée du modèle de données",
        "Figure 3 - Chaîne de recherche intelligente et réponse ancrée",
        "Figure 4 - Séquence simplifiée de la messagerie instantanée",
        "Figure 5 - Recommandation personnalisée explicable",
    ])

    add_major_title(doc, "Introduction générale")
    add_text(doc, "Le marché de la pièce automobile d'occasion répond à un double besoin : réduire le coût de l'entretien et prolonger la durée de vie de composants encore utilisables. Pourtant, la recherche reste difficile. Une même pièce peut porter plusieurs références, convenir à plusieurs véhicules et dépendre de critères précis tels que l'année, le moteur, le carburant ou la position sur le véhicule. Une annonce incomplète ou une interprétation approximative peut donc mener à un achat inutile.", style="Lead")
    add_text(doc, "Le projet E-biessa a été conçu pour structurer ce marché et simplifier la mise en relation locale. Il ne s'agit pas seulement d'un catalogue : la plateforme associe des fonctions de place de marché, de confiance, de communication et d'intelligence artificielle. L'acheteur décrit son besoin avec ses propres mots, le système extrait les contraintes, recherche les annonces sémantiquement proches, puis vérifie la compatibilité déclarée avant d'afficher une réponse.")
    add_text(doc, "La problématique générale est la suivante : comment construire une place de marché automobile intuitive et riche, tout en garantissant que l'intelligence artificielle ne présente jamais une compatibilité non vérifiée ? La réponse retenue repose sur une séparation nette des responsabilités. Les données transactionnelles et l'autorisation restent dans Spring Boot et MySQL. La recherche intelligente évolue dans un service Python isolé. Les règles déterministes ont toujours priorité sur la sortie du modèle.")
    add_text(doc, "Le premier chapitre présente le contexte, la problématique et la méthode. Le deuxième formalise les besoins. Le troisième détaille la conception. Le quatrième décrit la réalisation finale. Le cinquième expose la stratégie de tests et les résultats. Une conclusion synthétise les apports et les perspectives.")

    # Chapter 1
    add_chapter(doc, 1, "Présentation du projet", "Ce chapitre situe le besoin, les utilisateurs visés, les contraintes du domaine automobile et la démarche qui a conduit à la solution finale.")
    doc.add_paragraph("1.1 Contexte", style="Heading 2")
    add_text(doc, "En Tunisie, les échanges de pièces d'occasion passent souvent par des annonces généralistes, des groupes sociaux ou des contacts directs. Ces canaux favorisent la proximité, mais ils structurent peu les informations essentielles. Le titre d'une annonce peut mentionner un véhicule sans préciser la génération, le moteur ou la position. Les coordonnées du vendeur sont parfois exposées publiquement, tandis que les mécanismes de réputation et de modération restent limités.")
    add_text(doc, "E-biessa centralise ces échanges dans une application spécialisée. Le catalogue impose des champs cohérents, permet plusieurs véhicules compatibles et distingue la référence OEM principale des références interchangeables. Les fonctionnalités de confiance - validation des vendeurs, avis, signalements, suspension et journal d'audit - accompagnent la recherche et la transaction.")

    doc.add_paragraph("1.2 Étude de l'existant", style="Heading 2")
    add_table(doc, ["Approche", "Forces", "Limites observées"], [
        ("Petites annonces généralistes", "Audience large, publication rapide", "Compatibilité peu structurée, filtres insuffisants"),
        ("Groupes sociaux", "Échanges directs, communauté locale", "Recherche instable, données dispersées, modération variable"),
        ("Catalogues professionnels", "Références précises, stocks organisés", "Souvent centrés sur le neuf et sur un inventaire fermé"),
        ("E-biessa", "Domaine spécialisé, recherche hybride, confiance intégrée", "Dépend de la qualité des données saisies et validées"),
    ], widths=[3.8, 5.6, 6.4])

    doc.add_paragraph("1.3 Problématique", style="Heading 2")
    add_callout(doc, "Question directrice", "Comment offrir une recherche naturelle et rapide de pièces d'occasion, tout en maintenant une compatibilité vérifiable, une identité fiable des utilisateurs et des outils de modération adaptés ?", fill=ORANGE)
    add_text(doc, "Trois difficultés structurent le projet. Premièrement, les données véhicule doivent être propres et reliées à chaque annonce sans mélanger des caractéristiques provenant de véhicules différents. Deuxièmement, les rôles doivent limiter chaque action : un acheteur ne publie pas avant validation comme vendeur, un vendeur ne modère pas, et un administrateur dispose d'un historique d'actions. Troisièmement, l'intelligence artificielle doit améliorer la compréhension sans remplacer les règles métier.")

    doc.add_paragraph("1.4 Solution proposée", style="Heading 2")
    add_bullets(doc, [
        "une application Angular responsive, utilisable en thème clair ou sombre sans modification de géométrie ;",
        "un noyau Spring Boot sécurisé pour les comptes, annonces, images, messages, avis et opérations administratives ;",
        "un service FastAPI pour l'indexation, la recherche vectorielle et l'extraction de contraintes ;",
        "une base MySQL comme source de vérité et un index Qdrant reconstruit à partir des annonces actives ;",
        "un modèle Qwen adapté au vocabulaire français, anglais et tunisien, encadré par des validations déterministes ;",
        "un moteur de recommandations explicable fondé sur les préférences véhicule et les interactions privées ;",
        "une couche agentique en lecture seule pour orchestrer recherche, compatibilité et recommandation."
    ])

    doc.add_paragraph("1.5 Démarche de réalisation", style="Heading 2")
    add_text(doc, "La réalisation a suivi une démarche incrémentale proche de Scrum. Chaque incrément a produit une fonctionnalité démontrable, accompagnée de tests et d'une mise à jour de la documentation. Cette organisation a permis de stabiliser la marketplace avant d'introduire l'intelligence artificielle.")
    add_table(doc, ["Incrément", "Contenu principal", "Critère de sortie"], [
        ("Socle", "Authentification, rôles, profils, catalogue et annonces", "Parcours acheteur/vendeur cohérents"),
        ("Confiance", "Validation vendeur, avis, signalements, administration", "Actions auditées et autorisées"),
        ("Communication", "Contact protégé, messagerie et images", "Échange temps réel entre deux comptes"),
        ("IA", "Embeddings, Qdrant, Qwen adapté et recherche ancrée", "Résultats compatibles et explicables"),
        ("Personnalisation", "Préférences, historique minimal et agents", "Suggestions privées, contrôlables et justifiées"),
        ("Validation", "Tests unitaires, intégration et E2E", "Scénarios critiques reproductibles"),
    ], widths=[2.7, 7.5, 5.7])

    # Chapter 2
    add_chapter(doc, 2, "Analyse et spécification des besoins", "L'analyse transforme la problématique en fonctions mesurables et en contraintes de qualité applicables aux trois profils de la plateforme.")
    doc.add_paragraph("2.1 Identification des acteurs", style="Heading 2")
    add_table(doc, ["Acteur", "Responsabilités et objectifs"], [
        ("Visiteur", "Consulter le catalogue public, filtrer et ouvrir une annonce sans accéder aux coordonnées privées."),
        ("Acheteur", "Gérer son profil et ses véhicules, sauvegarder, contacter, échanger, signaler et évaluer un vendeur éligible."),
        ("Vendeur", "Gérer ses annonces, images, compatibilités et statuts, répondre aux conversations et suivre sa réputation."),
        ("Administrateur", "Valider les vendeurs, modérer annonces et signalements, suspendre des comptes et consulter l'audit."),
        ("Service IA", "Indexer les annonces, comprendre les requêtes et classer des candidats sans modifier les données métier."),
    ], widths=[3.0, 12.8])

    doc.add_paragraph("2.2 Besoins fonctionnels", style="Heading 2")
    add_bullets(doc, [
        "inscription par courriel et mot de passe, connexion OAuth facultative, récupération par courriel ou téléphone ;",
        "édition d'un profil complet avec coordonnées, localisation, activité, biographie et site web ;",
        "demande de passage vendeur placée en attente jusqu'à décision de l'administrateur ;",
        "cycle de vie complet d'une annonce : brouillon, aperçu, publication, modification, désactivation, vente et suppression ;",
        "ajout, compression, prévisualisation, ordre par glisser-déposer, choix de l'image principale et suppression ;",
        "recherche par catégorie, marque, modèle, année, carburant, moteur, position, prix, livraison, état et tri ;",
        "recherche en langage naturel avec contraintes comprises et raisons de correspondance ;",
        "contact protégé par téléphone, WhatsApp, courriel et messagerie interne avec pièces jointes ;",
        "avis de 1 à 5 étoiles, commentaire facultatif, réputation agrégée et signalement vendeur ;",
        "tableau de bord administrateur, files de traitement et journal d'audit ;",
        "recommandations personnalisées, paramètres de confidentialité et effacement de l'historique de recommandation."
    ])

    doc.add_paragraph("2.3 Besoins non fonctionnels", style="Heading 2")
    add_table(doc, ["Axe", "Exigence retenue"], [
        ("Sécurité", "JWT signé, mots de passe hachés, RBAC, secrets hors dépôt, validation de fichiers et limitation de débit."),
        ("Fiabilité", "MySQL reste la source de vérité ; un index vectoriel obsolète ne peut pas rendre une annonce inactive."),
        ("Explicabilité", "Toute recommandation et tout résultat IA affichent des raisons issues de données contrôlées."),
        ("Performance", "Pagination, cache court, seuil de pertinence, concurrence bornée et images compressées."),
        ("Ergonomie", "Interface responsive, états de chargement, erreurs lisibles, thème clair/sombre géométriquement stable."),
        ("Maintenabilité", "Séparation contrôleur/service/dépôt, DTO dédiés, configuration externalisée et tests automatisés."),
        ("Confidentialité", "Coordonnées réservées aux acheteurs connectés ; historique IA minimal et effaçable."),
        ("Portabilité", "Services exécutables localement sur Windows pour la démonstration académique."),
    ], widths=[3.0, 12.8])

    doc.add_paragraph("2.4 Règles métier centrales", style="Heading 2")
    add_numbered(doc, [
        "Une annonce publique doit être active, appartenir à un vendeur non suspendu et contenir des données valides.",
        "Une annonce marquée vendue disparaît immédiatement des listes, recherches et recommandations publiques.",
        "Les coordonnées du vendeur ne sont révélées qu'à un acheteur authentifié.",
        "Le passage au rôle vendeur exige une demande et une décision explicite de l'administrateur.",
        "Une compatibilité multi-véhicule est évaluée ligne par ligne ; le système ne combine jamais la marque d'une ligne avec le carburant d'une autre.",
        "La sortie du modèle de langage est une proposition de contraintes, jamais une preuve de compatibilité.",
        "Les messages et avis sont accessibles uniquement aux participants ou aux rôles autorisés.",
        "Chaque décision administrative sensible produit une entrée d'audit."
    ])

    doc.add_paragraph("2.5 Matrice de traçabilité", style="Heading 2")
    add_table(doc, ["Besoin", "Composant responsable", "Validation"], [
        ("Authentifier", "Spring Security, JWT, OAuth2", "Tests de rôles et parcours navigateur"),
        ("Publier une annonce", "PartService, Angular Seller Studio", "CRUD, DTO, upload et statuts"),
        ("Vérifier la compatibilité", "FastAPI + règles déterministes Spring", "Jeux multilingues et contradictions"),
        ("Communiquer", "MessagingService, WebSocket", "Conversation privée et pièces jointes"),
        ("Modérer", "AdminService, audit", "Files de signalement et suspension"),
        ("Recommander", "RecommendationService, agents en lecture", "Scénarios de préférences, exclusion et diversité"),
    ], widths=[3.4, 6.5, 5.9])

    # Chapter 3
    add_chapter(doc, 3, "Conception", "La conception sépare les responsabilités transactionnelles, l'expérience utilisateur et le calcul d'intelligence artificielle afin de réduire le couplage et de garder les décisions critiques vérifiables.")
    doc.add_paragraph("3.1 Choix d'architecture", style="Heading 2")
    add_text(doc, "Le cahier des charges initial mettait l'accent sur FastAPI. Or, le socle transactionnel du projet a été stabilisé en Spring Boot. La décision finale conserve Spring comme coeur métier et exploite FastAPI comme microservice IA. Cette adaptation respecte les capacités demandées tout en évitant une réécriture risquée.")
    add_callout(doc, "Décision structurante", "Spring Boot est la frontière de sécurité et la source des réponses publiques. FastAPI ne reçoit que les documents nécessaires à l'indexation et renvoie des identifiants classés. Avant affichage, Spring recharge toujours les annonces depuis MySQL.")
    add_figure(doc, ASSETS / "architecture.png", "Figure 1 - Architecture globale de la solution")

    doc.add_paragraph("3.2 Architecture logique", style="Heading 2")
    add_table(doc, ["Couche", "Responsabilité"], [
        ("Présentation", "Composants Angular, routes, formulaires, thèmes, retours d'erreur et accessibilité."),
        ("API", "Contrôleurs REST, validation des DTO, sérialisation stable et contrats HTTP."),
        ("Métier", "Services Spring portant autorisation, règles de cycle de vie, modération et classement."),
        ("Persistance", "Dépôts JPA et modèle relationnel MySQL ; Qdrant pour les vecteurs reconstruisibles."),
        ("IA", "Embeddings, extraction de contraintes, récupération, agents spécialisés et explications."),
        ("Infrastructure", "Stockage d'images, courriel/SMS optionnels, OAuth et catalogue véhicule externe avec repli local."),
    ], widths=[3.0, 12.8])

    doc.add_paragraph("3.3 Modèle de données", style="Heading 2")
    add_text(doc, "L'entité Utilisateur centralise l'identité et le rôle. Une annonce appartient à un vendeur, à une catégorie et à une collection de véhicules compatibles. Les images conservent un ordre et une image principale. Les conversations relient un acheteur, un vendeur et l'annonce d'origine. Les signaux d'activité restent séparés des messages et des coordonnées afin de limiter les données utilisées pour la personnalisation.")
    add_figure(doc, ASSETS / "data_model.png", "Figure 2 - Vue simplifiée du modèle de données", width_cm=15.5)

    doc.add_paragraph("3.4 Conception de la recherche intelligente", style="Heading 2")
    add_text(doc, "Une annonce active est transformée en document textuel normalisé contenant titre, description, catégorie, références et compatibilités. FastEmbed calcule son vecteur multilingue, stocké dans la collection Qdrant. Lors d'une requête, Qwen extrait un objet JSON. Le système fusionne cette sortie avec les contraintes lexicales et les filtres explicites, puis récupère des candidats sémantiques.")
    add_text(doc, "Le classement final applique un seuil minimal et vérifie les contraintes sur une même ligne de compatibilité. Les raisons visibles - par exemple « compatible Golf », « diesel », « position arrière » ou « système de freinage » - proviennent des champs réellement associés à l'annonce. La réponse est construite par Spring à partir des DTO actifs, ce qui empêche une prose générée ou un identifiant périmé de devenir une recommandation publique.")
    add_figure(doc, ASSETS / "ai_pipeline.png", "Figure 3 - Chaîne de recherche intelligente et réponse ancrée")

    doc.add_paragraph("3.5 Orchestration agentique finale", style="Heading 2")
    add_text(doc, "La couche LangChain coordonne trois outils spécialisés en lecture seule. L'agent de recherche traduit l'intention en requête de récupération. L'agent de compatibilité contrôle les contraintes structurées et explique les rejets. L'agent de recommandation utilise le flux personnalisé déjà calculé par Spring. Un orchestrateur choisit l'outil utile, agrège les résultats et remet à Spring des éléments structurés. Aucun agent ne crée, ne modifie ou ne supprime une annonce.")
    add_table(doc, ["Agent", "Entrées", "Sortie autorisée", "Garde-fou"], [
        ("Recherche", "Texte et filtres", "Identifiants candidats", "Limite et seuil de pertinence"),
        ("Compatibilité", "Contraintes et fiche annonce", "Compatible, incompatible ou inconnu", "Règles catalogue prioritaires"),
        ("Recommandation", "Identité authentifiée et limite", "Suggestions et raisons", "Pas de données inter-utilisateurs"),
        ("Orchestrateur", "Intention utilisateur", "Plan court et réponse structurée", "Lecture seule, délai et nombre d'appels bornés"),
    ], widths=[2.8, 4.0, 4.6, 4.4])

    doc.add_paragraph("3.6 Sécurité", style="Heading 2")
    add_bullets(doc, [
        "identité extraite du contexte JWT, jamais d'un identifiant fourni librement par le client ;",
        "autorisation par rôle au niveau des routes et réévaluation de la propriété dans les services ;",
        "codes OAuth courts, à usage unique, sans JWT dans l'URL de redirection ;",
        "codes de récupération expirables, nombre d'essais limité et réponses non révélatrices ;",
        "jetons WebSocket courts et à usage unique pour éviter de transmettre le JWT dans une chaîne de requête ;",
        "signature réelle des fichiers image contrôlée en plus du type MIME déclaré ;",
        "jeton interne partagé entre Spring et FastAPI pour les opérations d'indexation ;",
        "limitation de débit, concurrence bornée, cache court et délai maximum sur l'endpoint IA public ;",
        "secrets chargés depuis l'environnement et exclus de Git."
    ])

    doc.add_paragraph("3.7 Conception de la messagerie", style="Heading 2")
    add_text(doc, "La messagerie suit un modèle conversationnel unique par contexte d'annonce. Le chargement initial place la vue sur le dernier message, tandis que le contenu reste dans une zone de hauteur limitée et défilable. Les événements WebSocket mettent à jour l'interface sans recharger la page. Les images jointes sont stockées via le même contrat que les images d'annonces, avec des limites dédiées.")
    add_figure(doc, ASSETS / "messaging_sequence.png", "Figure 4 - Séquence simplifiée de la messagerie instantanée")

    doc.add_paragraph("3.8 Conception de la recommandation", style="Heading 2")
    add_text(doc, "Le moteur combine compatibilité déclarée, catégories sauvegardées, vendeurs contactés, popularité et récence. Les préférences véhicule agissent comme contrainte forte. Les interactions comportementales sont dédupliquées, limitées à 200 enregistrements, pondérées avec une demi-vie de trente jours et ignorées après quatre-vingt-dix jours. Le score reste un rang interne et n'est jamais présenté comme un pourcentage de compatibilité.")
    add_figure(doc, ASSETS / "recommendations.png", "Figure 5 - Recommandation personnalisée explicable")

    # Chapter 4
    add_chapter(doc, 4, "Réalisation", "Cette partie présente l'environnement technique, les principaux modules développés et l'intégration finale des services.")
    doc.add_paragraph("4.1 Environnement de travail", style="Heading 2")
    add_table(doc, ["Élément", "Choix final", "Rôle"], [
        ("Langage backend", "Java 21", "API transactionnelle et règles métier"),
        ("Framework backend", "Spring Boot 3.2.8", "Web, sécurité, validation, JPA et WebSocket"),
        ("Frontend", "Angular 21, TypeScript 5.9", "SPA responsive et formulaires réactifs"),
        ("Service IA", "Python 3.12, FastAPI", "Recherche sémantique et orchestration"),
        ("Données", "MySQL, Qdrant", "Source de vérité relationnelle et index vectoriel"),
        ("Modèle", "Qwen 2.5 7B Instruct", "Extraction multilingue de contraintes"),
        ("Adaptation", "Unsloth QLoRA", "Entraînement local quantifié"),
        ("Exécution LLM", "Ollama", "Inférence locale sur GPU"),
        ("Tests", "JUnit, Vitest, Pytest, Playwright", "Validation multi-couche"),
        ("Versionnement", "Git et GitHub", "Historique, partage et sauvegarde du code"),
    ], widths=[3.0, 4.4, 8.4])

    doc.add_paragraph("4.2 Backend Spring Boot", style="Heading 2")
    add_text(doc, "Le backend applique une organisation par contrôleurs, services, dépôts, entités et DTO. Les contrôleurs exposent les contrats HTTP et délèguent les décisions aux services. Les DTO empêchent la sérialisation directe des entités JPA et stabilisent les données visibles. Le gestionnaire global d'exceptions transforme les erreurs en messages exploitables par Angular.")
    add_text(doc, "Les domaines couverts sont l'authentification, les utilisateurs, les catégories, les véhicules, les annonces, les images, les favoris, les conversations, les avis, les signalements, les recommandations et l'administration. Les événements de synchronisation avec l'IA sont publiés uniquement après la validation de la transaction ; une panne de l'index n'annule donc pas une opération légitime sur le catalogue.")

    doc.add_paragraph("4.3 Frontend Angular", style="Heading 2")
    add_text(doc, "L'application Angular est organisée en vues autonomes : marketplace, détail d'annonce, favoris, profil, profil public du vendeur, studio vendeur, boîte de réception et tableau de bord administrateur. La navigation et les actions changent selon le rôle. Les formulaires signalent les champs invalides et affichent les messages d'erreur retournés par l'API.")
    add_text(doc, "Le thème clair utilise des surfaces légèrement teintées et des accents verts et orange pour délimiter les zones, tandis que le thème sombre conserve exactement les mêmes espacements, bordures et tailles. La bascule de thème modifie seulement les variables de couleur, ce qui supprime les glissements de mise en page.")
    add_callout(doc, "Principe d'interface", "Une action métier importante possède toujours un état visible : chargement, succès, erreur récupérable ou confirmation. Les listes sont paginées et les filtres restent dans l'URL afin de pouvoir revenir à la même recherche.")

    h44 = doc.add_paragraph("4.4 Gestion complète des annonces", style="Heading 2")
    h44.paragraph_format.page_break_before = True
    add_table(doc, ["Bloc", "Informations prises en charge"], [
        ("Identification", "Titre, description, catégorie, fabricant, marque de pièce, OEM principal et alternatives"),
        ("État commercial", "Prix, négociable, état, garantie, livraison, collecte et localisation"),
        ("Compatibilité", "Plusieurs véhicules, année, moteur, carburant et position"),
        ("Médias", "JPEG/PNG/WebP, compression, ordre, image principale et suppression"),
        ("Cycle de vie", "Aperçu, actif, inactif, vendu, modéré et supprimé"),
        ("Découverte", "Filtres, tri, pagination, favoris, recherche IA et recommandations"),
    ], widths=[3.4, 12.4])

    doc.add_paragraph("4.5 Comptes, profils et confiance", style="Heading 2")
    add_text(doc, "L'identité locale est liée à l'adresse électronique et peut être complétée par Google, Facebook ou Microsoft lorsque les identifiants OAuth correspondants sont configurés. Le profil séparé contient nom, téléphone, localisation, nom commercial, site web et biographie. L'acheteur qui souhaite vendre soumet ces informations ; son statut reste en attente jusqu'à la décision d'un administrateur.")
    add_text(doc, "Après une interaction admissible, l'acheteur attribue de une à cinq étoiles et peut ajouter un commentaire. La moyenne et le nombre d'avis apparaissent sur le profil public. Un signalement peut viser une annonce ou un vendeur. Le tableau de bord administrateur regroupe les demandes, rapports, utilisateurs, annonces et traces d'audit.")

    doc.add_paragraph("4.6 Contact et messagerie", style="Heading 2")
    add_text(doc, "La fiche annonce comporte une section « Contacter le vendeur ». Pour réduire l'extraction automatisée, les coordonnées ne sont renvoyées qu'à un acheteur connecté. Les boutons génèrent un appel téléphonique, un message WhatsApp préparé avec le titre de la pièce ou un courriel. L'acheteur peut aussi démarrer une conversation interne.")
    add_text(doc, "La boîte de réception présente une liste organisée de conversations et un panneau de discussion. Le fil possède une hauteur limitée, s'ouvre sur le dernier message et conserve le défilement vers les anciens échanges. L'envoi d'une image permet au vendeur de fournir des vues supplémentaires d'une pièce sans exposer un canal externe.")

    doc.add_paragraph("4.7 Mise en oeuvre de l'IA", style="Heading 2")
    add_text(doc, "Le modèle de base Qwen 2.5 7B Instruct a été chargé en 4 bits puis adapté avec Unsloth et LoRA de rang 8 sur un ordinateur équipé d'une RTX 4060 de 8 Go de VRAM et de 16 Go de mémoire système. Le jeu d'apprentissage contient des formulations anglaises, françaises et tunisiennes ainsi que les sorties JSON canoniques attendues. Une passe d'entraînement a été retenue après un test court, car les erreurs restantes concernaient surtout la copie de chiffres et relevaient mieux d'une validation déterministe que d'époques supplémentaires.")
    add_text(doc, "Le modèle exporté est servi localement par Ollama sous le nom e-biessa-query-v2. FastAPI lui demande uniquement un objet structuré. La fonction de fusion contrôle ensuite l'année explicite, normalise les catégories et vérifie chaque valeur dans le catalogue. Cette approche apporte la souplesse linguistique du LLM tout en gardant un comportement reproductible.")

    doc.add_paragraph("4.8 Jeu de démonstration", style="Heading 2")
    add_text(doc, "Le profil de développement initialise un administrateur, trois acheteurs, trois vendeurs et soixante annonces réparties entre plusieurs marques, modèles, catégories et villes. Chaque vendeur possède plusieurs annonces détaillées. Les données sont réutilisables sans duplication au redémarrage et servent uniquement à la démonstration académique.")
    add_table(doc, ["Donnée", "Volume de démonstration"], [
        ("Administrateurs", "1"),
        ("Acheteurs", "3"),
        ("Vendeurs", "3"),
        ("Annonces actives initiales", "60"),
        ("Définitions véhicule locales", "20"),
        ("Langues d'exemple IA", "anglais, français, arabe tunisien translittéré"),
    ], widths=[7.0, 8.8])

    doc.add_paragraph("4.9 Configuration et exécution locale", style="Heading 2")
    add_numbered(doc, [
        "Démarrer MySQL et créer la base marketplace.",
        "Définir les variables DB_URL, DB_USERNAME, DB_PASSWORD, JWT_SECRET et CORS_ALLOWED_ORIGINS dans l'environnement local de Spring.",
        "Définir le même AI_SERVICE_TOKEN dans Spring et FastAPI, puis activer l'intégration IA.",
        "Lancer Ollama avec le modèle adapté e-biessa-query-v2.",
        "Démarrer FastAPI sur 127.0.0.1:8001, Spring Boot sur localhost:8080 et Angular sur localhost:4200.",
        "Exécuter une réindexation complète avec le compte administrateur avant la première recherche sémantique."
    ])

    # Chapter 5
    add_chapter(doc, 5, "Tests et validation", "La validation combine des tests rapides par composant et des parcours représentatifs. Les chiffres présentés correspondent aux campagnes automatisées conservées avec le projet et à des jeux de données contrôlés.")
    doc.add_paragraph("5.1 Stratégie", style="Heading 2")
    add_text(doc, "La pyramide de tests couvre d'abord les règles métier isolées, puis les intégrations avec une base de test, les contrats du service IA et enfin les parcours dans un navigateur réel. Les scénarios sensibles - identité, propriété d'une annonce, compatibilité et visibilité des éléments vendus - sont testés à plusieurs niveaux.")
    add_table(doc, ["Niveau", "Outil", "Portée"], [
        ("Backend", "JUnit, Spring Boot Test, H2", "Services, contrôleurs, sécurité, sérialisation et persistance"),
        ("Frontend", "Vitest", "Composants, formulaires, thèmes, erreurs et recommandations"),
        ("IA", "Pytest", "Sécurité interne, normalisation, recherche et garde-fous"),
        ("Navigateur", "Playwright", "Rôles, recherche IA, thèmes et parcours principaux"),
        ("Modèle", "Jeu de validation tenu à l'écart", "JSON et exactitude par champ"),
    ], widths=[2.8, 4.2, 8.8])

    doc.add_paragraph("5.2 Résultats automatisés", style="Heading 2")
    add_table(doc, ["Suite", "Résultat validé", "Observation"], [
        ("Spring Boot", "50 tests réussis", "Inclut recommandations, sécurité, réponses IA et limitation de débit"),
        ("Angular", "44 tests réussis", "Inclut profil, interface, erreurs et recommandations"),
        ("FastAPI", "20 tests réussis", "Un avertissement de dépréciation non bloquant"),
        ("Playwright isolé", "5 parcours réussis", "Navigateur réel avec API contrôlée"),
        ("Playwright pile complète", "Scénarios opt-in", "Nécessite MySQL, Spring, FastAPI, Ollama et Angular actifs"),
    ], widths=[3.2, 4.0, 8.6])
    add_callout(doc, "Lecture correcte des résultats", "Les tests Playwright isolés prouvent le comportement de l'interface dans un navigateur, mais ne remplacent pas une campagne sur la pile locale complète. Les deux suites sont donc conservées séparément.", fill=ORANGE)

    doc.add_paragraph("5.3 Évaluation du modèle de compréhension", style="Heading 2")
    add_text(doc, "L'évaluation tenue à l'écart comporte 120 requêtes. Le modèle de base comprenait déjà correctement la marque, le modèle, l'année et le carburant, mais normalisait mal la catégorie et le type de pièce. L'adaptation QLoRA a nettement amélioré la sortie canonique.")
    add_table(doc, ["Mesure", "Résultat final"], [
        ("JSON valide", "99,2 % - 119/120"),
        ("Marque", "99,2 %"),
        ("Modèle", "96,7 %"),
        ("Carburant", "99,2 %"),
        ("Position", "99,2 %"),
        ("Catégorie canonique", "99,2 %"),
    ], widths=[9.2, 6.6])
    add_text(doc, "Après normalisation, le type de pièce canonique atteint lui aussi 99,2 %. Après validation autoritaire de l'année et du catalogue, l'exactitude d'un enregistrement complet atteint 96,7 % (116/120). Ces résultats justifient l'usage du modèle comme extracteur. Ils ne justifient pas de lui déléguer la décision de compatibilité. Les erreurs de sérialisation d'année sont corrigées par l'extraction lexicale et la validation catalogue.")

    doc.add_paragraph("5.4 Validation de la recherche et des recommandations", style="Heading 2")
    add_text(doc, "Le benchmark de récupération comprend six requêtes synthétiques liées au catalogue de démonstration. Le taux de présence de la bonne annonce dans les douze premiers résultats, le rang réciproque moyen et la précision des contraintes atteignent 1,0 sur ce petit jeu. Ces valeurs confirment la cohérence technique, mais ne constituent pas une mesure statistique de qualité en conditions réelles.")
    add_text(doc, "Les recommandations sont vérifiées sur le démarrage à froid, la correspondance exacte multi-véhicule, les carburants contradictoires, la priorité aux catégories sauvegardées, les vendeurs contactés, la diversité, les restrictions de rôle, l'effacement de l'historique et l'exclusion d'annonces supprimées. Les raisons retournées proviennent des mêmes signaux que le score.")

    doc.add_paragraph("5.5 Recette fonctionnelle", style="Heading 2")
    add_table(doc, ["Scénario", "Résultat attendu"], [
        ("Acheteur recherche un étrier avant pour Golf diesel 2020", "Contraintes reconnues, aucune pièce arrière affichée"),
        ("Vendeur marque une annonce vendue", "Disparition immédiate du catalogue et de l'index public"),
        ("Acheteur ouvre les coordonnées", "Accès accordé uniquement avec le rôle Acheteur actif"),
        ("Acheteur joint une photo à un message", "Fichier validé, stocké et diffusé aux deux participants"),
        ("Acheteur demande à devenir vendeur", "Statut en attente jusqu'à décision administrateur"),
        ("Administrateur suspend un vendeur", "Connexion bloquée, annonces masquées et action auditée"),
        ("Acheteur désactive la personnalisation", "Historique comportemental supprimé, préférences explicites conservées"),
        ("Index Qdrant contient un ID obsolète", "Spring l'écarte lors du rechargement MySQL"),
    ], widths=[7.3, 8.5])

    doc.add_paragraph("5.6 Limites et mesures compensatoires", style="Heading 2")
    add_table(doc, ["Limite", "Mesure intégrée"], [
        ("Catalogue CarQuery parfois indisponible", "Repli local et confirmation explicite de la compatibilité par le vendeur"),
        ("Petit jeu d'évaluation IA", "Résultats présentés comme techniques, sans généralisation abusive"),
        ("Données de démonstration", "Séparation des profils local et production, identifiants non réutilisables"),
        ("LLM sujet aux erreurs", "Sortie structurée, normalisation, seuil et règles déterministes"),
        ("Classement personnalisé content-based", "Explications, désactivation, effacement et préparation à des évaluations humaines"),
        ("Exécution académique locale", "Architecture documentée et configuration externe pour une future mise en ligne"),
    ], widths=[5.4, 10.4])

    # Conclusion
    add_major_title(doc, "Conclusion générale")
    add_text(doc, "Le projet E-biessa aboutit à une plateforme complète de mise en relation autour des pièces automobiles d'occasion. Le socle transactionnel prend en charge les comptes, les rôles, les profils, les annonces riches, les images, la compatibilité multi-véhicule, les favoris, le contact, la messagerie, la réputation et l'administration. L'expérience Angular rend ces fonctions accessibles dans une interface cohérente et responsive.", style="Lead")
    add_text(doc, "L'apport majeur est l'intégration contrôlée de l'intelligence artificielle. Le modèle adapté comprend des formulations multilingues proches du langage des utilisateurs. La recherche vectorielle améliore le rappel, tandis que la validation déterministe protège la précision. Les recommandations et les agents spécialisés utilisent les mêmes principes : lecture seule, données minimales, raisons visibles et retour systématique à la source de vérité.")
    add_text(doc, "Ce travail a permis de mobiliser la conception orientée domaine, la sécurité web, les architectures distribuées, l'UX, le traitement du langage et l'évaluation expérimentale. La séparation Spring Boot / FastAPI représente un compromis justifié entre robustesse métier et agilité IA.")
    doc.add_paragraph("Perspectives", style="Heading 2")
    add_bullets(doc, [
        "déployer la pile derrière HTTPS après rotation et gestion centralisée des secrets ;",
        "faire valider la compatibilité par des professionnels et intégrer des catalogues OEM sous licence ;",
        "constituer un benchmark plus large, annoté manuellement, couvrant dialectes et fautes d'orthographe ;",
        "mesurer la qualité des recommandations avec des retours anonymisés et des métriques hors ligne ;",
        "ajouter notifications, paiement ou réservation uniquement après validation juridique et opérationnelle ;",
        "surveiller dérive du modèle, qualité des données et taux de signalement avant tout usage public."
    ])

    add_major_title(doc, "Bibliographie et webographie")
    refs = [
        "[1] Spring, Spring Boot Reference Documentation, https://docs.spring.io/spring-boot/",
        "[2] Spring, Spring Security Reference, https://docs.spring.io/spring-security/reference/",
        "[3] Angular, Angular Documentation, https://angular.dev/",
        "[4] FastAPI, FastAPI Documentation, https://fastapi.tiangolo.com/",
        "[5] Qdrant, Qdrant Documentation, https://qdrant.tech/documentation/",
        "[6] Ollama, Ollama Documentation, https://docs.ollama.com/",
        "[7] Qwen Team, Qwen2.5 Technical Report, https://arxiv.org/abs/2412.15115",
        "[8] Unsloth AI, Unsloth Documentation, https://docs.unsloth.ai/",
        "[9] LangChain, LangChain Documentation, https://docs.langchain.com/",
        "[10] OWASP Foundation, OWASP Top 10, https://owasp.org/www-project-top-ten/",
        "[11] IETF, RFC 7519 - JSON Web Token, https://www.rfc-editor.org/rfc/rfc7519",
        "[12] CarQuery API, Vehicle Data API, https://www.carqueryapi.com/",
        "[13] Cloudinary, Image and Video API Documentation, https://cloudinary.com/documentation/",
        "[14] Microsoft, Playwright Documentation, https://playwright.dev/",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        r = p.add_run(ref)
        set_run_font(r, size=9.5)

    add_major_title(doc, "Annexe A - Cartographie des API")
    add_table(doc, ["Domaine", "Exemples d'opérations"], [
        ("Authentification", "register, login, OAuth exchange, recovery request/complete"),
        ("Annonces", "list, detail, create, update, status, delete, my-listings"),
        ("Images", "upload, primary, order, delete"),
        ("Vendeurs", "profile, listings, contact, reviews, reports"),
        ("Messagerie", "inbox, start, detail, send text/image, realtime ticket"),
        ("Recommandations", "feed, settings, update settings, clear history"),
        ("IA", "semantic search, full reindex"),
        ("Administration", "users, seller applications, reports, moderation, audit"),
    ], widths=[4.0, 11.8])

    doc.add_paragraph("Annexe B - Variables de configuration", style="Heading 2")
    add_table(doc, ["Variable", "Usage"], [
        ("DB_URL / DB_USERNAME / DB_PASSWORD", "Connexion MySQL"),
        ("JWT_SECRET", "Signature des jetons d'accès"),
        ("CORS_ALLOWED_ORIGINS", "Origines frontend autorisées"),
        ("AI_SERVICE_TOKEN", "Authentification interne Spring/FastAPI"),
        ("AI_SERVICE_ENABLED", "Activation de la recherche sémantique"),
        ("OLLAMA_BASE_URL / OLLAMA_MODEL", "Accès au modèle local"),
        ("CLOUDINARY_*", "Stockage distant optionnel des images"),
        ("GOOGLE_* / FACEBOOK_* / MICROSOFT_*", "Fournisseurs OAuth optionnels"),
        ("SMTP_* / TWILIO_*", "Récupération de compte hors profil de développement"),
    ], widths=[6.0, 9.8])
    add_callout(doc, "Règle de gestion", "Les fichiers .env, les uploads, les index Qdrant, les caches, les modèles GGUF, les adaptateurs LoRA et les résultats Playwright sont exclus du dépôt Git.")

    doc.add_paragraph("Annexe C - Commandes de validation", style="Heading 2")
    commands = [
        ("Backend", "cd marketplace-api\n.\\mvnw.cmd test"),
        ("Frontend", "cd marketplace-frontend\nnpm run verify"),
        ("Service IA", "cd ai-service\n.\\.venv\\Scripts\\python.exe -m pytest"),
        ("E2E isolé", "cd marketplace-frontend\nnpm run test:e2e"),
        ("E2E pile complète", "cd marketplace-frontend\nnpm run test:e2e:live"),
    ]
    for name, cmd in commands:
        p = doc.add_paragraph()
        r = p.add_run(name)
        set_run_font(r, bold=True, color=GREEN)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.left_indent = Cm(0.8)
        for idx, line in enumerate(cmd.splitlines()):
            rr = p2.add_run(line)
            rr.font.name = "Consolas"
            rr._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            rr.font.size = Pt(9)
            if idx < len(cmd.splitlines()) - 1:
                rr.add_break()

    # Document properties and automatic field update.
    doc.core_properties.title = "Conception et réalisation de E-biessa"
    doc.core_properties.subject = "Rapport de stage académique - marketplace intelligente de pièces automobiles d'occasion"
    doc.core_properties.author = "Jasser Noomani"
    doc.core_properties.keywords = "E-biessa, Spring Boot, Angular, FastAPI, Qwen, Qdrant, marketplace automobile"
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
